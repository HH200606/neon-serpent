# -*- coding: utf-8 -*-
"""把 docs/PRD.md 转成排版过的 Word 文档。

只覆盖这份 PRD 用到的 Markdown 子集：标题、表格、列表、引用、代码块、
行内加粗与行内代码。目的是可复现地产出对外传阅版本，而不是通用转换器。
"""

from __future__ import annotations

import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = None  # 由 main() 注入
OUT = None  # 由 main() 注入

INK = RGBColor(0x26, 0x22, 0x20)
INK_SOFT = RGBColor(0x5C, 0x54, 0x4C)
MUTED = RGBColor(0x8C, 0x82, 0x76)
CINNABAR = RGBColor(0x9E, 0x3A, 0x2E)

TITLE = None
FOOTER = None  # 缺省时从 frontmatter 的 version 推导，避免页脚和正文版本漂移

LATIN = "Microsoft YaHei"
EASTASIA = "微软雅黑"
BODY_EASTASIA = "宋体"
MONO = "Consolas"


def set_eastasia(run, name: str):
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), name)
    fonts.set(qn("w:ascii"), LATIN)
    fonts.set(qn("w:hAnsi"), LATIN)


def style_fonts(doc):
    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_EASTASIA)

    for name, size, color in (
        ("Title", 20, INK),
        ("Heading 1", 16, INK),
        ("Heading 2", 13, INK),
        ("Heading 3", 11.5, INK_SOFT),
    ):
        st = doc.styles[name]
        st.font.name = LATIN
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.element.rPr.rFonts.set(qn("w:eastAsia"), EASTASIA)


def add_page_footer(doc):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{FOOTER or '产品需求文档'}　·　第 ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED
    set_eastasia(run, BODY_EASTASIA)

    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)

    run2 = p.add_run(" 页")
    run2.font.size = Pt(8.5)
    run2.font.color.rgb = MUTED
    set_eastasia(run2, BODY_EASTASIA)


def shade(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
URL = re.compile(r"https?://[^\s)>\]]+")
ZWSP = "\u200b"


def soften_urls(text: str) -> str:
    """在 URL 内部插入零宽空格，让 Word 能在版心内折行（否则会顶破页边距）。"""

    def repl(m):
        url = m.group(0)
        out = []
        for ch in url:
            out.append(ch)
            if ch in "/-._?=&%":
                out.append(ZWSP)
        return "".join(out)

    return URL.sub(repl, text)


def set_mono(run):
    """代码片段：等宽字体，同时给中文指定回退字体，避免拉丁设置把中文也改掉。"""
    run.font.name = MONO
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:ascii"), MONO)
    fonts.set(qn("w:hAnsi"), MONO)
    fonts.set(qn("w:eastAsia"), BODY_EASTASIA)


def add_inline(paragraph, text: str, size=None, color=None):
    text = soften_urls(text)
    for part in INLINE.split(text):
        if not part:
            continue
        is_code = part.startswith("`") and part.endswith("`")
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif is_code:
            run = paragraph.add_run(part[1:-1])
            set_mono(run)
            run.font.size = Pt((size or 10.5) - 1)
        else:
            run = paragraph.add_run(part)
        if size:
            run.font.size = Pt(size)
        if not is_code:
            if color:
                run.font.color.rgb = color
            set_eastasia(run, BODY_EASTASIA)
        elif color:
            run.font.color.rgb = color


def parse_table_row(line: str):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


TEXT_WIDTH_CM = 16.2  # A4（21cm）减去左右各 2.4cm 页边距


def column_widths(count: int):
    """按列数给一组固定的列宽（cm），保证表格不超出版心。"""
    presets = {
        2: [3.6, 11.4],
        3: [2.4, 6.3, 6.3],
        4: [2.0, 5.2, 4.8, 3.0],
        5: [1.5, 3.7, 3.7, 3.1, 3.0],
        6: [1.3, 3.0, 3.0, 2.7, 2.6, 2.4],
    }
    if count in presets:
        return presets[count]
    each = TEXT_WIDTH_CM / count
    return [each] * count


def fix_table_width(table):
    """把表格锁成「版心 100% 宽 + 自动布局」，避免列宽算法把表顶出页边距。"""
    tblPr = table._tbl.tblPr
    for el in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(el)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:type"), "pct")
    tblW.set(qn("w:w"), "5000")



def read_frontmatter(lines):
    """把文档开头的 YAML 头读成扁平 dict（只用于取 version 等简单字段）。"""
    data = {}
    if lines and lines[0].strip() == "---":
        for j in range(1, len(lines)):
            raw = lines[j]
            if raw.strip() == "---":
                break
            if ":" in raw:
                key, value = raw.split(":", 1)
                data[key.strip()] = value.strip()
    return data


def strip_frontmatter(lines):
    """跳过文档开头的 YAML 头（给 agent 读的那部分不进 Word）。"""
    if lines and lines[0].strip() == "---":
        for j in range(1, len(lines)):
            if lines[j].strip() == "---":
                return lines[j + 1:]
    return lines


def collect_toc(lines):
    return [ln.strip()[3:].strip() for ln in lines if ln.strip().startswith("## ")]


def add_toc(doc, entries):
    if not entries:
        return
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    head = doc.add_paragraph()
    hr = head.add_run("目录")
    hr.bold = True
    hr.font.size = Pt(13)
    set_eastasia(hr, EASTASIA)
    head.paragraph_format.space_before = Pt(10)
    head.paragraph_format.space_after = Pt(6)
    for e in entries:
        q = doc.add_paragraph()
        q.paragraph_format.left_indent = Cm(0.4)
        q.paragraph_format.space_after = Pt(2)
        rr = q.add_run(e)
        rr.font.size = Pt(9.5)
        rr.font.color.rgb = INK_SOFT
        set_eastasia(rr, BODY_EASTASIA)
    brk = doc.add_paragraph()
    brk.add_run().add_break(WD_BREAK.PAGE)


def add_image(doc, alt, path):
    if not os.path.exists(path):
        warn = doc.add_paragraph()
        wr = warn.add_run("[缺图] " + alt + "：" + path)
        wr.font.size = Pt(9)
        wr.font.color.rgb = MUTED
        set_eastasia(wr, BODY_EASTASIA)
        return
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.space_before = Pt(6)
    pic.paragraph_format.space_after = Pt(2)
    pic.add_run().add_picture(path, width=Cm(14))
    if alt:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(10)
        cr = cap.add_run(alt)
        cr.font.size = Pt(8.5)
        cr.font.color.rgb = MUTED
        set_eastasia(cr, BODY_EASTASIA)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    fix_table_width(table)

    hdr = table.rows[0].cells
    for i, text in enumerate(header):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        add_inline(p, text, size=8.5)
        for run in p.runs:
            run.bold = True
        shade(hdr[i], "EFE5D2")

    for row in body:
        cells = table.add_row().cells
        for i, text in enumerate(row[: len(header)]):
            cells[i].text = ""
            add_inline(cells[i].paragraphs[0], text, size=8.5)

    doc.add_paragraph()


def main(argv=None):
    import argparse
    global SRC, OUT, TITLE, FOOTER
    ap = argparse.ArgumentParser(description="把 PRD.md 转成排版好的 docx")
    ap.add_argument("--src", default=os.path.join(ROOT, "docs", "PRD.md"))
    ap.add_argument("--out", default=os.path.join(ROOT, "docs", "PRD.docx"))
    ap.add_argument("--title", default=None, help="封面标题；缺省取一级标题")
    ap.add_argument("--footer", default=None, help="页脚前缀；缺省按 frontmatter 的 version 生成")
    a = ap.parse_args(argv)
    SRC, OUT, TITLE, FOOTER = a.src, a.out, a.title, a.footer
    build()


def build():
    global FOOTER
    with open(SRC, encoding="utf-8") as fh:
        lines = fh.read().splitlines()

    meta = read_frontmatter(lines)
    lines = strip_frontmatter(lines)
    TOC = collect_toc(lines)

    if FOOTER is None:
        version = meta.get("version", "")
        short = ".".join(version.split(".")[:2]) if version else ""
        FOOTER = f"产品需求文档 v{short}" if short else "产品需求文档"

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)

    style_fonts(doc)
    add_page_footer(doc)

    i = 0
    first_h1 = True
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 代码块
        if stripped.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run("\n".join(block))
            run.font.name = MONO
            run.font.size = Pt(8.5)
            run.font.color.rgb = INK_SOFT
            continue

        # 目录占位
        if stripped == "<!--toc-->":
            add_toc(doc, TOC)
            i += 1
            continue

        # 图片
        if stripped.startswith("!["):
            m = re.match(r"^!\[(.*?)\]\((.*?)\)$", stripped)
            if m:
                add_image(doc, m.group(1), os.path.join(os.path.dirname(SRC), m.group(2)))
                i += 1
                continue

        # 表格
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = parse_table_row(lines[i])
                if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in row):
                    rows.append(row)
                i += 1
            if rows:
                add_table(doc, rows)
            continue

        if stripped == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        if stripped.startswith("# "):
            text = stripped[2:]
            if first_h1:
                p = doc.add_paragraph(style="Title")
                add_inline(p, TITLE or text)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                first_h1 = False
            else:
                add_inline(doc.add_heading("", level=1), text)
            i += 1
            continue

        if stripped.startswith("## "):
            add_inline(doc.add_heading("", level=1), stripped[3:])
            i += 1
            continue

        if stripped.startswith("### "):
            add_inline(doc.add_heading("", level=2), stripped[4:])
            i += 1
            continue

        if stripped.startswith("#### "):
            add_inline(doc.add_heading("", level=3), stripped[5:])
            i += 1
            continue

        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, stripped[2:], color=INK_SOFT)
            i += 1
            continue

        if re.match(r"^[-*]\s+", stripped):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^[-*]\s+", "", stripped))
            i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s+", "", stripped))
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        add_inline(p, stripped)
        i += 1

    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
